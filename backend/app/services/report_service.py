from __future__ import annotations

import json
import logging
import math
import statistics
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from uuid import UUID, uuid4

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


logger = logging.getLogger(__name__)

BACKEND_ROOT = Path(__file__).resolve().parents[2]
REPORT_ROOT = BACKEND_ROOT / "data" / "reports"
DATASET_ROOT = BACKEND_ROOT / "data" / "datasets"
PREDICTION_ROOT = BACKEND_ROOT / "data" / "predictions"
STREAM_ROOT = BACKEND_ROOT / "data" / "streams"
ALERT_ROOT = BACKEND_ROOT / "data" / "alerts"

REPORT_ROOT.mkdir(parents=True, exist_ok=True)

SUPPORTED_FORMATS = ("pdf", "xlsx", "csv", "json", "html")
DEFAULT_FORMATS = ("pdf", "xlsx", "json", "html")
MAX_TABLE_ROWS = 5_000

PREDICTION_COLUMNS = {
    "predicted_lithology",
    "predicted_porosity",
    "predicted_water_saturation",
    "predicted_permeability_md",
    "hydrocarbon_probability",
    "reservoir_quality",
    "lithology_confidence",
}

DEPTH_CANDIDATES = (
    "depth_m",
    "depth",
    "DEPTH",
    "tvd_m",
    "tvdss_m",
)

THIN_GREY = Side(style="thin", color="D9E2F3")
HEADER_FILL = "1F4E78"
SUBHEADER_FILL = "D9EAF7"
ACCENT_FILL = "E2F0D9"
WARNING_FILL = "FFF2CC"
CRITICAL_FILL = "F4CCCC"


# ---------------------------------------------------------------------------
# Dataclasses
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class ReportRequest:
    source_type: str
    source_id: str
    title: str | None = None
    well_id: str | None = None
    formats: tuple[str, ...] = DEFAULT_FORMATS
    include_alerts: bool = True
    include_raw_records: bool = False
    include_recommendations: bool = True
    reservoir_porosity_cutoff: float = 0.10
    reservoir_sw_cutoff: float = 0.60
    hydrocarbon_probability_cutoff: float = 0.60
    minimum_interval_thickness: float = 0.5
    metadata: Mapping[str, Any] | None = None


@dataclass
class ReservoirInterval:
    top_depth: float
    base_depth: float
    gross_thickness: float
    sample_count: int
    mean_porosity: float | None
    mean_water_saturation: float | None
    mean_permeability_md: float | None
    mean_hydrocarbon_probability: float | None
    dominant_lithology: str | None
    reservoir_quality: str | None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class ReportAnalysis:
    executive_summary: dict[str, Any]
    data_quality: dict[str, Any]
    petrophysical_summary: dict[str, Any]
    reservoir_intervals: list[ReservoirInterval]
    alert_summary: dict[str, Any]
    model_summary: dict[str, Any]
    recommendations: list[str]

    def to_dict(self) -> dict[str, Any]:
        return {
            "executive_summary": self.executive_summary,
            "data_quality": self.data_quality,
            "petrophysical_summary": self.petrophysical_summary,
            "reservoir_intervals": [
                item.to_dict()
                for item in self.reservoir_intervals
            ],
            "alert_summary": self.alert_summary,
            "model_summary": self.model_summary,
            "recommendations": list(self.recommendations),
        }


@dataclass
class ReportResult:
    report_id: str
    created_at: str
    source_type: str
    source_id: str
    title: str
    well_id: str | None
    report_directory: str
    generated_files: dict[str, str]
    analysis: ReportAnalysis
    metadata: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return {
            "report_id": self.report_id,
            "created_at": self.created_at,
            "source_type": self.source_type,
            "source_id": self.source_id,
            "title": self.title,
            "well_id": self.well_id,
            "report_directory": self.report_directory,
            "generated_files": self.generated_files,
            "analysis": self.analysis.to_dict(),
            "metadata": self.metadata,
        }


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------

def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_json_value(value: Any) -> Any:
    if isinstance(value, np.generic):
        value = value.item()

    if isinstance(value, float):
        return value if math.isfinite(value) else None

    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()

    if isinstance(value, Path):
        return str(value)

    if isinstance(value, UUID):
        return str(value)

    try:
        if pd.isna(value):
            return None
    except Exception:
        pass

    return value


def _json_default(value: Any) -> Any:
    safe = _safe_json_value(value)
    if safe is not value:
        return safe

    if isinstance(value, np.ndarray):
        return value.tolist()

    if isinstance(value, set):
        return sorted(value)

    raise TypeError(
        f"Object of type {type(value).__name__} "
        "is not JSON serialisable."
    )


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
            allow_nan=False,
            default=_json_default,
        ),
        encoding="utf-8",
    )
    temporary.replace(path)


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}

    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload if isinstance(payload, dict) else {}


def _read_jsonl(path: Path, limit: int | None = None) -> list[dict[str, Any]]:
    if not path.exists():
        return []

    records: list[dict[str, Any]] = []

    with path.open("r", encoding="utf-8") as source:
        for line in source:
            if not line.strip():
                continue
            payload = json.loads(line)
            if isinstance(payload, dict):
                records.append(payload)

    if limit is not None:
        records = records[-limit:]

    return records


def _normalise_formats(formats: Sequence[str]) -> tuple[str, ...]:
    normalised = tuple(
        dict.fromkeys(
            str(item).strip().lower()
            for item in formats
            if str(item).strip()
        )
    )

    unknown = [
        item
        for item in normalised
        if item not in SUPPORTED_FORMATS
    ]

    if unknown:
        raise ValueError(
            "Unsupported report formats: "
            + ", ".join(unknown)
        )

    return normalised or DEFAULT_FORMATS


def _validate_uuid(identifier: str, label: str) -> None:
    try:
        UUID(identifier)
    except ValueError as exc:
        raise ValueError(f"Invalid {label}: {identifier}") from exc


def _coerce_numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce")


def _first_existing_column(
    dataframe: pd.DataFrame,
    candidates: Iterable[str],
) -> str | None:
    for candidate in candidates:
        if candidate in dataframe.columns:
            return candidate
    return None


def _find_well_id(
    dataframe: pd.DataFrame,
    metadata: Mapping[str, Any],
    requested_well_id: str | None,
) -> str | None:
    if requested_well_id:
        return requested_well_id

    for key in ("well_id", "Well_ID", "well", "well_name"):
        value = metadata.get(key)
        if value:
            return str(value)

    for column in ("well_id", "Well_ID", "well", "well_name"):
        if column in dataframe.columns:
            non_null = dataframe[column].dropna()
            if not non_null.empty:
                return str(non_null.iloc[0])

    return None


def _round_or_none(value: Any, digits: int = 4) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None

    if not math.isfinite(number):
        return None

    return round(number, digits)


def _series_stats(series: pd.Series) -> dict[str, Any]:
    numeric = _coerce_numeric(series).dropna()

    if numeric.empty:
        return {
            "count": 0,
            "minimum": None,
            "maximum": None,
            "mean": None,
            "median": None,
            "standard_deviation": None,
            "p10": None,
            "p90": None,
        }

    return {
        "count": int(numeric.count()),
        "minimum": _round_or_none(numeric.min()),
        "maximum": _round_or_none(numeric.max()),
        "mean": _round_or_none(numeric.mean()),
        "median": _round_or_none(numeric.median()),
        "standard_deviation": _round_or_none(numeric.std(ddof=0)),
        "p10": _round_or_none(numeric.quantile(0.10)),
        "p90": _round_or_none(numeric.quantile(0.90)),
    }


def _categorical_distribution(series: pd.Series) -> dict[str, int]:
    if series.empty:
        return {}

    counts = (
        series.fillna("Unknown")
        .astype(str)
        .value_counts(dropna=False)
    )

    return {
        str(index): int(value)
        for index, value in counts.items()
    }


# ---------------------------------------------------------------------------
# Source loading
# ---------------------------------------------------------------------------

class ReportSourceLoader:
    def load(
        self,
        source_type: str,
        source_id: str,
    ) -> tuple[pd.DataFrame, dict[str, Any]]:
        source_type = source_type.strip().lower()

        if source_type == "dataset":
            return self._load_dataset(source_id)

        if source_type == "prediction":
            return self._load_prediction(source_id)

        if source_type == "stream":
            return self._load_stream(source_id)

        raise ValueError(
            "source_type must be one of: "
            "dataset, prediction, stream"
        )

    def _load_dataset(
        self,
        dataset_id: str,
    ) -> tuple[pd.DataFrame, dict[str, Any]]:
        # Platform V1 datasets use human-readable ds-* identifiers and are
        # registered in SQLite. Legacy UUID dataset directories remain supported.
        if dataset_id.startswith("ds-"):
            from app.platform_v1.database import connection
            with connection() as conn:
                row = conn.execute("SELECT * FROM datasets WHERE dataset_id = ?", (dataset_id,)).fetchone()
            if row is None:
                raise FileNotFoundError(f"Dataset not found: {dataset_id}")
            path = Path(row["file_path"])
            if not path.exists():
                raise FileNotFoundError(f"Dataset file was not found: {path}")
            if path.suffix.lower() == ".csv": dataframe = pd.read_csv(path)
            elif path.suffix.lower() == ".parquet": dataframe = pd.read_parquet(path)
            elif path.suffix.lower() == ".las":
                import lasio
                dataframe = lasio.read(path).df().reset_index()
            else: raise ValueError("Unsupported dataset format for reporting.")
            metadata = {"dataset_id": dataset_id, "name": row["name"], "file_name": row["file_name"], "row_count": row["row_count"], "columns": json.loads(row["columns_json"])}
            return dataframe, metadata
        _validate_uuid(dataset_id, "dataset ID")
        directory = DATASET_ROOT / dataset_id
        records_path = directory / "processed" / "records.csv"
        if not directory.exists(): raise FileNotFoundError(f"Dataset not found: {dataset_id}")
        if not records_path.exists(): raise FileNotFoundError("Processed dataset records were not found.")
        dataframe = pd.read_csv(records_path)
        metadata = _read_json(directory / "metadata.json")
        metadata.setdefault("dataset_id", dataset_id)
        return dataframe, metadata

    def _load_prediction(
        self,
        prediction_id: str,
    ) -> tuple[pd.DataFrame, dict[str, Any]]:
        _validate_uuid(prediction_id, "prediction ID")
        directory = PREDICTION_ROOT / prediction_id
        records_path = directory / "predictions.csv"

        if not directory.exists():
            raise FileNotFoundError(
                f"Prediction result not found: {prediction_id}"
            )

        if not records_path.exists():
            raise FileNotFoundError(
                "Prediction CSV was not found."
            )

        dataframe = pd.read_csv(records_path)
        metadata = _read_json(directory / "metadata.json")
        metadata.setdefault("prediction_id", prediction_id)
        return dataframe, metadata

    def _load_stream(
        self,
        stream_id: str,
    ) -> tuple[pd.DataFrame, dict[str, Any]]:
        _validate_uuid(stream_id, "stream ID")
        directory = STREAM_ROOT / stream_id
        records_path = directory / "predictions.jsonl"

        if not directory.exists():
            raise FileNotFoundError(
                f"Stream session not found: {stream_id}"
            )

        records = _read_jsonl(records_path, limit=100_000)
        dataframe = pd.DataFrame.from_records(records)
        metadata = _read_json(directory / "metadata.json")
        metadata.setdefault("stream_id", stream_id)
        return dataframe, metadata


# ---------------------------------------------------------------------------
# Analysis
# ---------------------------------------------------------------------------

class ReportAnalyser:
    def analyse(
        self,
        dataframe: pd.DataFrame,
        *,
        request: ReportRequest,
        metadata: Mapping[str, Any],
        alerts: Sequence[Mapping[str, Any]],
    ) -> ReportAnalysis:
        well_id = _find_well_id(
            dataframe,
            metadata,
            request.well_id,
        )

        executive = self._executive_summary(
            dataframe,
            request=request,
            metadata=metadata,
            well_id=well_id,
        )

        data_quality = self._data_quality_summary(
            dataframe,
            metadata,
        )

        petrophysical = self._petrophysical_summary(
            dataframe
        )

        intervals = self._identify_intervals(
            dataframe,
            porosity_cutoff=request.reservoir_porosity_cutoff,
            sw_cutoff=request.reservoir_sw_cutoff,
            hydrocarbon_cutoff=request.hydrocarbon_probability_cutoff,
            minimum_thickness=request.minimum_interval_thickness,
        )

        alert_summary = self._alert_summary(alerts)

        model_summary = self._model_summary(metadata)

        recommendations = (
            self._recommendations(
                dataframe=dataframe,
                data_quality=data_quality,
                intervals=intervals,
                alert_summary=alert_summary,
                model_summary=model_summary,
            )
            if request.include_recommendations
            else []
        )

        return ReportAnalysis(
            executive_summary=executive,
            data_quality=data_quality,
            petrophysical_summary=petrophysical,
            reservoir_intervals=intervals,
            alert_summary=alert_summary,
            model_summary=model_summary,
            recommendations=recommendations,
        )

    def _executive_summary(
        self,
        dataframe: pd.DataFrame,
        *,
        request: ReportRequest,
        metadata: Mapping[str, Any],
        well_id: str | None,
    ) -> dict[str, Any]:
        depth_column = _first_existing_column(
            dataframe,
            DEPTH_CANDIDATES,
        )

        depth_min = None
        depth_max = None

        if depth_column:
            depth = _coerce_numeric(dataframe[depth_column]).dropna()
            if not depth.empty:
                depth_min = _round_or_none(depth.min(), 3)
                depth_max = _round_or_none(depth.max(), 3)

        qc_score = (
            metadata.get("qc_score")
            or metadata.get("quality_score")
            or metadata.get("ai_readiness_score")
        )

        ai_ready = metadata.get("ai_ready")
        if ai_ready is None:
            ai_ready = metadata.get("ai_readiness")

        return {
            "report_date": _utc_now(),
            "source_type": request.source_type,
            "source_id": request.source_id,
            "well_id": well_id,
            "sample_count": int(len(dataframe)),
            "column_count": int(len(dataframe.columns)),
            "depth_column": depth_column,
            "top_depth": depth_min,
            "base_depth": depth_max,
            "gross_depth_range": (
                _round_or_none(depth_max - depth_min, 3)
                if depth_min is not None and depth_max is not None
                else None
            ),
            "qc_score": _round_or_none(qc_score, 2),
            "ai_ready": ai_ready,
            "processing_time_seconds": metadata.get(
                "processing_time_seconds"
            ),
            "prediction_runtime_seconds": metadata.get(
                "runtime_seconds"
            ),
        }

    def _data_quality_summary(
        self,
        dataframe: pd.DataFrame,
        metadata: Mapping[str, Any],
    ) -> dict[str, Any]:
        missing_by_column = {
            str(column): int(dataframe[column].isna().sum())
            for column in dataframe.columns
            if int(dataframe[column].isna().sum()) > 0
        }

        duplicate_rows = int(dataframe.duplicated().sum())

        depth_column = _first_existing_column(
            dataframe,
            DEPTH_CANDIDATES,
        )
        duplicate_depths = 0
        irregular_sampling = None

        if depth_column:
            depth = _coerce_numeric(dataframe[depth_column]).dropna()
            duplicate_depths = int(depth.duplicated().sum())

            if len(depth) > 2:
                differences = depth.sort_values().diff().dropna()
                if not differences.empty:
                    median_step = float(differences.median())
                    if median_step != 0:
                        coefficient = (
                            float(differences.std(ddof=0))
                            / abs(median_step)
                        )
                        irregular_sampling = round(coefficient, 4)

        qc_before = metadata.get("qc_before", {})
        qc_after = metadata.get("qc_after", {})
        warnings = metadata.get("warnings", [])
        provenance = metadata.get("provenance", [])

        return {
            "missing_cells_total": int(dataframe.isna().sum().sum()),
            "missing_by_column": missing_by_column,
            "duplicate_rows": duplicate_rows,
            "duplicate_depths": duplicate_depths,
            "irregular_sampling_coefficient": irregular_sampling,
            "qc_before": qc_before,
            "qc_after": qc_after,
            "warnings": warnings,
            "provenance": provenance,
            "available_columns": [
                str(column)
                for column in dataframe.columns
            ],
        }

    def _petrophysical_summary(
        self,
        dataframe: pd.DataFrame,
    ) -> dict[str, Any]:
        result: dict[str, Any] = {}

        numeric_fields = {
            "porosity": "predicted_porosity",
            "water_saturation": "predicted_water_saturation",
            "permeability_md": "predicted_permeability_md",
            "hydrocarbon_probability": "hydrocarbon_probability",
            "lithology_confidence": "lithology_confidence",
        }

        for label, column in numeric_fields.items():
            result[label] = (
                _series_stats(dataframe[column])
                if column in dataframe.columns
                else _series_stats(pd.Series(dtype=float))
            )

        categorical_fields = {
            "lithology_distribution": "predicted_lithology",
            "reservoir_quality_distribution": "reservoir_quality",
        }

        for label, column in categorical_fields.items():
            result[label] = (
                _categorical_distribution(dataframe[column])
                if column in dataframe.columns
                else {}
            )

        return result

    def _identify_intervals(
        self,
        dataframe: pd.DataFrame,
        *,
        porosity_cutoff: float,
        sw_cutoff: float,
        hydrocarbon_cutoff: float,
        minimum_thickness: float,
    ) -> list[ReservoirInterval]:
        depth_column = _first_existing_column(
            dataframe,
            DEPTH_CANDIDATES,
        )

        required = {
            "predicted_porosity",
            "predicted_water_saturation",
            "hydrocarbon_probability",
        }

        if depth_column is None or not required.issubset(dataframe.columns):
            return []

        working = dataframe.copy()
        working["_depth"] = _coerce_numeric(working[depth_column])
        working["_phi"] = _coerce_numeric(
            working["predicted_porosity"]
        )
        working["_sw"] = _coerce_numeric(
            working["predicted_water_saturation"]
        )
        working["_hc"] = _coerce_numeric(
            working["hydrocarbon_probability"]
        )

        working = working.dropna(
            subset=["_depth", "_phi", "_sw", "_hc"]
        ).sort_values("_depth")

        if working.empty:
            return []

        working["_qualifies"] = (
            (working["_phi"] >= porosity_cutoff)
            & (working["_sw"] <= sw_cutoff)
            & (working["_hc"] >= hydrocarbon_cutoff)
        )

        depth_diffs = working["_depth"].diff().dropna()
        median_step = (
            float(depth_diffs.median())
            if not depth_diffs.empty
            else 0.5
        )
        allowed_gap = max(abs(median_step) * 1.5, minimum_thickness)

        intervals: list[ReservoirInterval] = []
        current_indices: list[int] = []
        previous_depth: float | None = None

        def flush(indices: list[int]) -> None:
            if not indices:
                return

            subset = working.loc[indices]
            top = float(subset["_depth"].min())
            base = float(subset["_depth"].max())
            thickness = max(
                base - top + abs(median_step),
                abs(median_step),
            )

            if thickness < minimum_thickness:
                return

            lithology = None
            quality = None

            if "predicted_lithology" in subset.columns:
                values = subset["predicted_lithology"].dropna().astype(str)
                if not values.empty:
                    lithology = values.mode().iloc[0]

            if "reservoir_quality" in subset.columns:
                values = subset["reservoir_quality"].dropna().astype(str)
                if not values.empty:
                    quality = values.mode().iloc[0]

            permeability = (
                _coerce_numeric(
                    subset["predicted_permeability_md"]
                ).mean()
                if "predicted_permeability_md" in subset.columns
                else None
            )

            intervals.append(
                ReservoirInterval(
                    top_depth=round(top, 3),
                    base_depth=round(base, 3),
                    gross_thickness=round(thickness, 3),
                    sample_count=int(len(subset)),
                    mean_porosity=_round_or_none(subset["_phi"].mean()),
                    mean_water_saturation=_round_or_none(subset["_sw"].mean()),
                    mean_permeability_md=_round_or_none(permeability),
                    mean_hydrocarbon_probability=_round_or_none(
                        subset["_hc"].mean()
                    ),
                    dominant_lithology=lithology,
                    reservoir_quality=quality,
                )
            )

        for index, row in working.iterrows():
            depth = float(row["_depth"])
            qualifies = bool(row["_qualifies"])

            if not qualifies:
                flush(current_indices)
                current_indices = []
                previous_depth = None
                continue

            if (
                previous_depth is not None
                and abs(depth - previous_depth) > allowed_gap
            ):
                flush(current_indices)
                current_indices = []

            current_indices.append(index)
            previous_depth = depth

        flush(current_indices)

        return intervals

    def _alert_summary(
        self,
        alerts: Sequence[Mapping[str, Any]],
    ) -> dict[str, Any]:
        by_severity = {
            "critical": 0,
            "warning": 0,
            "info": 0,
        }
        by_status: dict[str, int] = {}
        by_type: dict[str, int] = {}

        for alert in alerts:
            severity = str(alert.get("severity", "info")).lower()
            by_severity[severity] = by_severity.get(severity, 0) + 1

            status_value = str(alert.get("status", "open")).lower()
            by_status[status_value] = by_status.get(status_value, 0) + 1

            alert_type = str(
                alert.get("alert_type", "unknown")
            )
            by_type[alert_type] = by_type.get(alert_type, 0) + 1

        return {
            "total_alerts": int(len(alerts)),
            "by_severity": by_severity,
            "by_status": by_status,
            "by_type": by_type,
            "alerts": [
                {
                    "alert_id": alert.get("alert_id"),
                    "rule_name": alert.get("rule_name"),
                    "severity": alert.get("severity"),
                    "message": alert.get("message"),
                    "recommendation": alert.get("recommendation"),
                    "depth_m": alert.get("depth_m"),
                    "triggered_at": alert.get("triggered_at"),
                }
                for alert in alerts[:250]
            ],
        }

    def _model_summary(
        self,
        metadata: Mapping[str, Any],
    ) -> dict[str, Any]:
        model = metadata.get("model", {})
        models = metadata.get("models", {})
        inference = metadata.get("inference", {})

        return {
            "model": model,
            "models": models,
            "inference": inference,
            "model_version": metadata.get("model_version"),
            "framework": metadata.get("framework"),
            "runtime_seconds": metadata.get("runtime_seconds"),
            "features_used": metadata.get("features_used", []),
            "missing_features": metadata.get("missing_features", []),
            "fallbacks_used": metadata.get("fallbacks_used", []),
        }

    def _recommendations(
        self,
        *,
        dataframe: pd.DataFrame,
        data_quality: Mapping[str, Any],
        intervals: Sequence[ReservoirInterval],
        alert_summary: Mapping[str, Any],
        model_summary: Mapping[str, Any],
    ) -> list[str]:
        recommendations: list[str] = []

        missing = data_quality.get("missing_by_column", {})
        if missing:
            recommendations.append(
                "Review missing values and confirm that the required "
                "logging curves were mapped correctly before final interpretation."
            )

        if int(data_quality.get("duplicate_depths", 0)) > 0:
            recommendations.append(
                "Resolve duplicate depth samples to prevent biased interval statistics."
            )

        irregular = data_quality.get("irregular_sampling_coefficient")
        if isinstance(irregular, (int, float)) and irregular > 0.25:
            recommendations.append(
                "Resample the logs onto a consistent depth grid before model retraining "
                "or high-resolution interval comparison."
            )

        if intervals:
            recommendations.append(
                f"Prioritise the {len(intervals)} identified prospective reservoir "
                "interval(s) for petrophysical validation and completion screening."
            )
        else:
            recommendations.append(
                "No interval met all configured pay-screening criteria; review the "
                "cut-offs and validate whether the available curves are sufficient."
            )

        critical = (
            alert_summary.get("by_severity", {})
            .get("critical", 0)
        )
        if critical:
            recommendations.append(
                f"Investigate the {critical} critical alert(s) before using the "
                "results for operational decisions."
            )

        fallbacks = model_summary.get("fallbacks_used", [])
        if fallbacks:
            recommendations.append(
                "One or more deterministic fallback models were used. Confirm model "
                "registry availability and validate fallback outputs against reference data."
            )

        if "predicted_water_saturation" in dataframe.columns:
            sw = _coerce_numeric(
                dataframe["predicted_water_saturation"]
            ).dropna()
            if not sw.empty and float(sw.mean()) > 0.70:
                recommendations.append(
                    "Average predicted water saturation is high; review possible water-bearing "
                    "intervals and confirm resistivity and Rw assumptions."
                )

        if "hydrocarbon_probability" in dataframe.columns:
            probability = _coerce_numeric(
                dataframe["hydrocarbon_probability"]
            ).dropna()
            if not probability.empty and float(probability.max()) >= 0.80:
                recommendations.append(
                    "High-probability hydrocarbon intervals should be cross-checked with "
                    "porosity, saturation, lithology and borehole-condition indicators."
                )

        return recommendations


# ---------------------------------------------------------------------------
# Alert loading
# ---------------------------------------------------------------------------

def _load_related_alerts(
    source_type: str,
    source_id: str,
    well_id: str | None,
) -> list[dict[str, Any]]:
    path = ALERT_ROOT / "alerts.jsonl"

    if not path.exists():
        return []

    alerts = _read_jsonl(path)
    matched: list[dict[str, Any]] = []

    for alert in alerts:
        matches_source = (
            alert.get("source_id") == source_id
            or alert.get("dataset_id") == source_id
            or alert.get("stream_id") == source_id
            or alert.get("prediction_id") == source_id
        )

        matches_well = (
            well_id is not None
            and str(alert.get("well_id")) == str(well_id)
        )

        if matches_source or matches_well:
            matched.append(alert)

    matched.sort(
        key=lambda item: str(item.get("triggered_at", "")),
        reverse=True,
    )
    return matched


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

class JsonReportWriter:
    def write(
        self,
        path: Path,
        result: ReportResult,
        dataframe: pd.DataFrame,
        *,
        include_raw_records: bool,
    ) -> None:
        payload = result.to_dict()

        if include_raw_records:
            safe = dataframe.replace([np.inf, -np.inf], np.nan)
            safe = safe.astype(object).where(pd.notna(safe), None)
            payload["records"] = safe.to_dict(orient="records")

        _write_json(path, payload)


class CsvReportWriter:
    def write(
        self,
        path: Path,
        dataframe: pd.DataFrame,
    ) -> None:
        dataframe.to_csv(path, index=False)


class HtmlReportWriter:
    def write(
        self,
        path: Path,
        result: ReportResult,
    ) -> None:
        analysis = result.analysis
        executive = analysis.executive_summary
        quality = analysis.data_quality
        petro = analysis.petrophysical_summary
        intervals = analysis.reservoir_intervals
        alerts = analysis.alert_summary
        model = analysis.model_summary

        def esc(value: Any) -> str:
            import html
            return html.escape("" if value is None else str(value))

        def key_value_table(payload: Mapping[str, Any]) -> str:
            rows = []
            for key, value in payload.items():
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                rows.append(
                    f"<tr><th>{esc(key.replace('_', ' ').title())}</th>"
                    f"<td>{esc(value)}</td></tr>"
                )
            return "<table>" + "".join(rows) + "</table>"

        interval_rows = "".join(
            "<tr>"
            f"<td>{item.top_depth}</td>"
            f"<td>{item.base_depth}</td>"
            f"<td>{item.gross_thickness}</td>"
            f"<td>{esc(item.mean_porosity)}</td>"
            f"<td>{esc(item.mean_water_saturation)}</td>"
            f"<td>{esc(item.mean_permeability_md)}</td>"
            f"<td>{esc(item.mean_hydrocarbon_probability)}</td>"
            f"<td>{esc(item.dominant_lithology)}</td>"
            f"<td>{esc(item.reservoir_quality)}</td>"
            "</tr>"
            for item in intervals
        )

        recommendation_items = "".join(
            f"<li>{esc(item)}</li>"
            for item in analysis.recommendations
        )

        html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(result.title)}</title>
<style>
body {{
    font-family: Arial, Helvetica, sans-serif;
    margin: 32px;
    color: #1f2937;
    line-height: 1.5;
}}
h1, h2 {{
    color: #1F4E78;
}}
h1 {{
    border-bottom: 3px solid #1F4E78;
    padding-bottom: 10px;
}}
section {{
    margin: 28px 0;
}}
table {{
    width: 100%;
    border-collapse: collapse;
    margin: 12px 0;
    font-size: 13px;
}}
th, td {{
    border: 1px solid #d1d5db;
    padding: 8px;
    vertical-align: top;
}}
th {{
    background: #D9EAF7;
    text-align: left;
}}
.badge {{
    display: inline-block;
    background: #E2F0D9;
    padding: 4px 8px;
    border-radius: 4px;
    margin-right: 8px;
}}
.footer {{
    margin-top: 40px;
    font-size: 12px;
    color: #6b7280;
}}
</style>
</head>
<body>
<h1>{esc(result.title)}</h1>
<p>
<span class="badge">Report ID: {esc(result.report_id)}</span>
<span class="badge">Source: {esc(result.source_type)}</span>
<span class="badge">Well: {esc(result.well_id or "Not specified")}</span>
</p>

<section>
<h2>Executive Summary</h2>
{key_value_table(executive)}
</section>

<section>
<h2>Data Quality</h2>
{key_value_table(quality)}
</section>

<section>
<h2>Petrophysical Summary</h2>
{key_value_table(petro)}
</section>

<section>
<h2>Reservoir Intervals</h2>
<table>
<thead>
<tr>
<th>Top Depth</th><th>Base Depth</th><th>Thickness</th>
<th>Mean Porosity</th><th>Mean Sw</th><th>Mean Permeability</th>
<th>Mean HC Probability</th><th>Lithology</th><th>Quality</th>
</tr>
</thead>
<tbody>{interval_rows}</tbody>
</table>
</section>

<section>
<h2>Alert Summary</h2>
{key_value_table(alerts)}
</section>

<section>
<h2>AI Model Summary</h2>
{key_value_table(model)}
</section>

<section>
<h2>Recommendations</h2>
<ul>{recommendation_items}</ul>
</section>

<div class="footer">
Generated by PetroEdge AI on {esc(result.created_at)}.
</div>
</body>
</html>
"""
        path.write_text(html_content, encoding="utf-8")


class ExcelReportWriter:
    def write(
        self,
        path: Path,
        result: ReportResult,
        dataframe: pd.DataFrame,
    ) -> None:
        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)

        self._write_summary_sheet(workbook, result)
        self._write_intervals_sheet(workbook, result.analysis.reservoir_intervals)
        self._write_alerts_sheet(
            workbook,
            result.analysis.alert_summary.get("alerts", []),
        )
        self._write_data_sheet(workbook, dataframe)
        self._write_metadata_sheet(workbook, result.metadata)

        workbook.save(path)

    def _title(self, sheet, text: str, end_column: int = 6) -> None:
        sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=end_column,
        )
        cell = sheet.cell(1, 1, text)
        cell.font = Font(
            bold=True,
            size=16,
            color="FFFFFF",
        )
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center")

    def _write_mapping(
        self,
        sheet,
        start_row: int,
        title: str,
        payload: Mapping[str, Any],
    ) -> int:
        sheet.cell(start_row, 1, title)
        sheet.cell(start_row, 1).font = Font(
            bold=True,
            color="FFFFFF",
        )
        sheet.cell(start_row, 1).fill = PatternFill(
            "solid",
            fgColor=HEADER_FILL,
        )
        sheet.merge_cells(
            start_row=start_row,
            start_column=1,
            end_row=start_row,
            end_column=4,
        )

        row = start_row + 1

        for key, value in payload.items():
            sheet.cell(row, 1, key.replace("_", " ").title())
            sheet.cell(row, 1).font = Font(bold=True)
            sheet.cell(row, 1).fill = PatternFill(
                "solid",
                fgColor=SUBHEADER_FILL,
            )

            if isinstance(value, (dict, list)):
                value = json.dumps(
                    value,
                    ensure_ascii=False,
                    default=_json_default,
                )

            sheet.cell(row, 2, value)
            sheet.merge_cells(
                start_row=row,
                start_column=2,
                end_row=row,
                end_column=4,
            )
            sheet.cell(row, 2).alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )
            row += 1

        return row + 1

    def _write_summary_sheet(
        self,
        workbook: Workbook,
        result: ReportResult,
    ) -> None:
        sheet = workbook.create_sheet("Summary")
        self._title(sheet, result.title, end_column=4)

        row = 3
        row = self._write_mapping(
            sheet,
            row,
            "Executive Summary",
            result.analysis.executive_summary,
        )
        row = self._write_mapping(
            sheet,
            row,
            "Data Quality",
            result.analysis.data_quality,
        )
        row = self._write_mapping(
            sheet,
            row,
            "Petrophysical Summary",
            result.analysis.petrophysical_summary,
        )
        row = self._write_mapping(
            sheet,
            row,
            "Model Summary",
            result.analysis.model_summary,
        )

        sheet.cell(row, 1, "Recommendations")
        sheet.cell(row, 1).font = Font(
            bold=True,
            color="FFFFFF",
        )
        sheet.cell(row, 1).fill = PatternFill(
            "solid",
            fgColor=HEADER_FILL,
        )
        sheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=4,
        )
        row += 1

        for item in result.analysis.recommendations:
            sheet.cell(row, 1, "•")
            sheet.cell(row, 2, item)
            sheet.merge_cells(
                start_row=row,
                start_column=2,
                end_row=row,
                end_column=4,
            )
            sheet.cell(row, 2).alignment = Alignment(wrap_text=True)
            row += 1

        widths = {1: 28, 2: 35, 3: 20, 4: 20}
        for index, width in widths.items():
            sheet.column_dimensions[get_column_letter(index)].width = width

        sheet.freeze_panes = "A3"

    def _write_intervals_sheet(
        self,
        workbook: Workbook,
        intervals: Sequence[ReservoirInterval],
    ) -> None:
        sheet = workbook.create_sheet("Reservoir Intervals")
        headers = [
            "Top Depth",
            "Base Depth",
            "Gross Thickness",
            "Sample Count",
            "Mean Porosity",
            "Mean Water Saturation",
            "Mean Permeability (mD)",
            "Mean Hydrocarbon Probability",
            "Dominant Lithology",
            "Reservoir Quality",
        ]

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(1, column, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        for row_index, interval in enumerate(intervals, start=2):
            values = [
                interval.top_depth,
                interval.base_depth,
                interval.gross_thickness,
                interval.sample_count,
                interval.mean_porosity,
                interval.mean_water_saturation,
                interval.mean_permeability_md,
                interval.mean_hydrocarbon_probability,
                interval.dominant_lithology,
                interval.reservoir_quality,
            ]

            for column, value in enumerate(values, start=1):
                sheet.cell(row_index, column, value)

        for column in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 18

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    def _write_alerts_sheet(
        self,
        workbook: Workbook,
        alerts: Sequence[Mapping[str, Any]],
    ) -> None:
        sheet = workbook.create_sheet("Alerts")
        headers = [
            "Alert ID",
            "Rule",
            "Severity",
            "Message",
            "Recommendation",
            "Depth",
            "Triggered At",
        ]

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(1, column, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)

        for row_index, alert in enumerate(alerts, start=2):
            values = [
                alert.get("alert_id"),
                alert.get("rule_name"),
                alert.get("severity"),
                alert.get("message"),
                alert.get("recommendation"),
                alert.get("depth_m"),
                alert.get("triggered_at"),
            ]

            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, column, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")

            severity = str(alert.get("severity", "")).lower()
            if severity == "critical":
                fill = PatternFill("solid", fgColor=CRITICAL_FILL)
            elif severity == "warning":
                fill = PatternFill("solid", fgColor=WARNING_FILL)
            else:
                fill = PatternFill("solid", fgColor=ACCENT_FILL)

            sheet.cell(row_index, 3).fill = fill

        widths = [38, 28, 12, 45, 45, 12, 24]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    def _write_data_sheet(
        self,
        workbook: Workbook,
        dataframe: pd.DataFrame,
    ) -> None:
        sheet = workbook.create_sheet("Data")

        limited = dataframe.head(MAX_TABLE_ROWS)

        for column_index, column in enumerate(limited.columns, start=1):
            cell = sheet.cell(1, column_index, str(column))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        for row_index, row in enumerate(
            limited.itertuples(index=False, name=None),
            start=2,
        ):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row_index, column_index, _safe_json_value(value))

        for column_index, column in enumerate(limited.columns, start=1):
            max_length = max(
                len(str(column)),
                max(
                    (
                        len(str(value))
                        for value in limited[column].head(200)
                        if pd.notna(value)
                    ),
                    default=0,
                ),
            )
            sheet.column_dimensions[
                get_column_letter(column_index)
            ].width = min(max(max_length + 2, 10), 32)

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    def _write_metadata_sheet(
        self,
        workbook: Workbook,
        metadata: Mapping[str, Any],
    ) -> None:
        sheet = workbook.create_sheet("Metadata")
        sheet.append(["Key", "Value"])

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)

        for key, value in metadata.items():
            if isinstance(value, (dict, list)):
                value = json.dumps(
                    value,
                    ensure_ascii=False,
                    default=_json_default,
                )
            sheet.append([key, value])

        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 90
        sheet.freeze_panes = "A2"


class PdfReportWriter:
    def write(
        self,
        path: Path,
        result: ReportResult,
    ) -> None:
        document = SimpleDocTemplate(
            str(path),
            pagesize=landscape(A4),
            rightMargin=16 * mm,
            leftMargin=16 * mm,
            topMargin=16 * mm,
            bottomMargin=16 * mm,
            title=result.title,
            author="PetroEdge AI",
        )

        styles = getSampleStyleSheet()
        styles.add(
            ParagraphStyle(
                name="PetroTitle",
                parent=styles["Title"],
                fontName="Helvetica-Bold",
                fontSize=20,
                leading=24,
                textColor=colors.HexColor("#1F4E78"),
                alignment=TA_CENTER,
                spaceAfter=14,
            )
        )
        styles.add(
            ParagraphStyle(
                name="PetroHeading",
                parent=styles["Heading2"],
                fontName="Helvetica-Bold",
                fontSize=13,
                leading=16,
                textColor=colors.HexColor("#1F4E78"),
                spaceBefore=10,
                spaceAfter=8,
            )
        )
        styles.add(
            ParagraphStyle(
                name="PetroBody",
                parent=styles["BodyText"],
                fontName="Helvetica",
                fontSize=8.5,
                leading=11,
                alignment=TA_LEFT,
            )
        )

        logo_path = BACKEND_ROOT / "app" / "assets" / "petroedge-logo.png"
        story: list[Any] = []
        if logo_path.exists():
            story.extend([Image(str(logo_path), width=72 * mm, height=40.5 * mm), Spacer(1, 5)])
        story.extend([
            Paragraph(result.title, styles["PetroTitle"]),
            Paragraph(
                f"Report ID: {result.report_id}<br/>"
                f"Source: {result.source_type} / {result.source_id}<br/>"
                f"Well: {result.well_id or 'Not specified'}<br/>"
                f"Generated: {result.created_at}",
                styles["PetroBody"],
            ),
            Spacer(1, 8),
        ])

        story.extend(
            self._mapping_section(
                "Executive Summary",
                result.analysis.executive_summary,
                styles,
            )
        )
        story.extend(
            self._mapping_section(
                "Data Quality",
                result.analysis.data_quality,
                styles,
            )
        )
        story.extend(
            self._mapping_section(
                "Petrophysical Summary",
                result.analysis.petrophysical_summary,
                styles,
            )
        )

        story.append(
            Paragraph("Reservoir Intervals", styles["PetroHeading"])
        )
        story.append(
            self._interval_table(
                result.analysis.reservoir_intervals,
                styles,
            )
        )

        story.extend(
            self._mapping_section(
                "Alert Summary",
                {
                    key: value
                    for key, value
                    in result.analysis.alert_summary.items()
                    if key != "alerts"
                },
                styles,
            )
        )

        story.append(
            Paragraph("Alert Details", styles["PetroHeading"])
        )
        story.append(
            self._alert_table(
                result.analysis.alert_summary.get("alerts", []),
                styles,
            )
        )

        story.extend(
            self._mapping_section(
                "AI Model Summary",
                result.analysis.model_summary,
                styles,
            )
        )

        story.append(
            Paragraph("Recommendations", styles["PetroHeading"])
        )

        if result.analysis.recommendations:
            for item in result.analysis.recommendations:
                story.append(
                    Paragraph(
                        f"• {item}",
                        styles["PetroBody"],
                    )
                )
                story.append(Spacer(1, 3))
        else:
            story.append(
                Paragraph(
                    "No recommendations were generated.",
                    styles["PetroBody"],
                )
            )

        document.build(
            story,
            onFirstPage=self._page_footer,
            onLaterPages=self._page_footer,
        )

    def _page_footer(self, canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(
            16 * mm,
            8 * mm,
            "PetroEdge AI - Confidential technical report",
        )
        canvas.drawRightString(
            landscape(A4)[0] - 16 * mm,
            8 * mm,
            f"Page {document.page}",
        )
        canvas.restoreState()

    def _mapping_section(
        self,
        title: str,
        payload: Mapping[str, Any],
        styles,
    ) -> list[Any]:
        rows = [["Metric", "Value"]]

        for key, value in payload.items():
            if isinstance(value, (dict, list)):
                value = json.dumps(
                    value,
                    ensure_ascii=False,
                    default=_json_default,
                )

            rows.append(
                [
                    Paragraph(
                        key.replace("_", " ").title(),
                        styles["PetroBody"],
                    ),
                    Paragraph(
                        str(value if value is not None else ""),
                        styles["PetroBody"],
                    ),
                ]
            )

        table = Table(
            rows,
            colWidths=[65 * mm, 190 * mm],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B7C9D6")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                        colors.white,
                        colors.HexColor("#F5F9FC"),
                    ]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )

        return [
            Paragraph(title, styles["PetroHeading"]),
            table,
            Spacer(1, 8),
        ]

    def _interval_table(
        self,
        intervals: Sequence[ReservoirInterval],
        styles,
    ) -> Table:
        headers = [
            "Top",
            "Base",
            "Thickness",
            "Samples",
            "Phi",
            "Sw",
            "Perm (mD)",
            "HC Prob.",
            "Lithology",
            "Quality",
        ]
        rows: list[list[Any]] = [headers]

        for item in intervals:
            rows.append(
                [
                    item.top_depth,
                    item.base_depth,
                    item.gross_thickness,
                    item.sample_count,
                    item.mean_porosity,
                    item.mean_water_saturation,
                    item.mean_permeability_md,
                    item.mean_hydrocarbon_probability,
                    item.dominant_lithology,
                    item.reservoir_quality,
                ]
            )

        if len(rows) == 1:
            rows.append(["No qualifying intervals", "", "", "", "", "", "", "", "", ""])

        table = Table(
            rows,
            colWidths=[
                22 * mm,
                22 * mm,
                24 * mm,
                19 * mm,
                18 * mm,
                18 * mm,
                24 * mm,
                23 * mm,
                35 * mm,
                32 * mm,
            ],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B7C9D6")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (0, 0), (7, -1), "CENTER"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                        colors.white,
                        colors.HexColor("#F5F9FC"),
                    ]),
                ]
            )
        )
        return table

    def _alert_table(
        self,
        alerts: Sequence[Mapping[str, Any]],
        styles,
    ) -> Table:
        rows: list[list[Any]] = [
            ["Severity", "Rule", "Depth", "Message", "Recommendation"]
        ]

        for alert in alerts[:100]:
            rows.append(
                [
                    alert.get("severity"),
                    alert.get("rule_name"),
                    alert.get("depth_m"),
                    Paragraph(
                        str(alert.get("message") or ""),
                        styles["PetroBody"],
                    ),
                    Paragraph(
                        str(alert.get("recommendation") or ""),
                        styles["PetroBody"],
                    ),
                ]
            )

        if len(rows) == 1:
            rows.append(["", "", "", "No related alerts.", ""])

        table = Table(
            rows,
            colWidths=[
                22 * mm,
                40 * mm,
                20 * mm,
                85 * mm,
                95 * mm,
            ],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B7C9D6")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                        colors.white,
                        colors.HexColor("#F5F9FC"),
                    ]),
                ]
            )
        )
        return table


# ---------------------------------------------------------------------------
# Report service
# ---------------------------------------------------------------------------

class ReportService:
    def __init__(
        self,
        *,
        loader: ReportSourceLoader | None = None,
        analyser: ReportAnalyser | None = None,
    ) -> None:
        self.loader = loader or ReportSourceLoader()
        self.analyser = analyser or ReportAnalyser()
        self.json_writer = JsonReportWriter()
        self.csv_writer = CsvReportWriter()
        self.html_writer = HtmlReportWriter()
        self.excel_writer = ExcelReportWriter()
        self.pdf_writer = PdfReportWriter()

    def generate(
        self,
        request: ReportRequest,
    ) -> ReportResult:
        formats = _normalise_formats(request.formats)

        dataframe, source_metadata = self.loader.load(
            request.source_type,
            request.source_id,
        )

        combined_metadata = {
            **source_metadata,
            **dict(request.metadata or {}),
        }

        well_id = _find_well_id(
            dataframe,
            combined_metadata,
            request.well_id,
        )

        alerts = (
            _load_related_alerts(
                request.source_type,
                request.source_id,
                well_id,
            )
            if request.include_alerts
            else []
        )

        analysis = self.analyser.analyse(
            dataframe,
            request=request,
            metadata=combined_metadata,
            alerts=alerts,
        )

        report_id = str(uuid4())
        created_at = _utc_now()
        title = (
            request.title
            or f"PetroEdge AI {request.source_type.title()} Report"
        )

        directory = REPORT_ROOT / report_id
        directory.mkdir(parents=True, exist_ok=False)

        result = ReportResult(
            report_id=report_id,
            created_at=created_at,
            source_type=request.source_type,
            source_id=request.source_id,
            title=title,
            well_id=well_id,
            report_directory=str(directory),
            generated_files={},
            analysis=analysis,
            metadata={
                **combined_metadata,
                "requested_formats": list(formats),
                "include_alerts": request.include_alerts,
                "include_raw_records": request.include_raw_records,
                "reservoir_porosity_cutoff": request.reservoir_porosity_cutoff,
                "reservoir_sw_cutoff": request.reservoir_sw_cutoff,
                "hydrocarbon_probability_cutoff": (
                    request.hydrocarbon_probability_cutoff
                ),
                "minimum_interval_thickness": request.minimum_interval_thickness,
            },
        )

        writers = {
            "json": lambda path: self.json_writer.write(
                path,
                result,
                dataframe,
                include_raw_records=request.include_raw_records,
            ),
            "csv": lambda path: self.csv_writer.write(path, dataframe),
            "html": lambda path: self.html_writer.write(path, result),
            "xlsx": lambda path: self.excel_writer.write(path, result, dataframe),
            "pdf": lambda path: self.pdf_writer.write(path, result),
        }

        filenames = {
            "json": "report.json",
            "csv": "report.csv",
            "html": "report.html",
            "xlsx": "report.xlsx",
            "pdf": "report.pdf",
        }

        for format_name in formats:
            output_path = directory / filenames[format_name]
            try:
                writers[format_name](output_path)
            except Exception:
                logger.exception(
                    "Report generation failed for format %s",
                    format_name,
                )
                raise

            result.generated_files[format_name] = str(output_path)

        metadata_path = directory / "metadata.json"
        result.generated_files["metadata"] = str(metadata_path)
        _write_json(metadata_path, result.to_dict())

        return result

    def list_reports(self) -> list[dict[str, Any]]:
        reports: list[dict[str, Any]] = []

        for directory in REPORT_ROOT.iterdir():
            if not directory.is_dir():
                continue

            metadata_path = directory / "metadata.json"
            if not metadata_path.exists():
                continue

            try:
                reports.append(_read_json(metadata_path))
            except Exception:
                logger.exception(
                    "Could not read report metadata from %s",
                    metadata_path,
                )

        reports.sort(
            key=lambda item: str(item.get("created_at", "")),
            reverse=True,
        )
        return reports

    def get_report(self, report_id: str) -> dict[str, Any]:
        _validate_uuid(report_id, "report ID")
        metadata_path = REPORT_ROOT / report_id / "metadata.json"

        if not metadata_path.exists():
            raise FileNotFoundError(
                f"Report not found: {report_id}"
            )

        return _read_json(metadata_path)

    def delete_report(self, report_id: str) -> bool:
        _validate_uuid(report_id, "report ID")
        directory = REPORT_ROOT / report_id

        if not directory.exists():
            return False

        import shutil

        shutil.rmtree(directory)
        return True


_report_service: ReportService | None = None


def get_report_service() -> ReportService:
    global _report_service

    if _report_service is None:
        _report_service = ReportService()

    return _report_service


def generate_report(request: ReportRequest) -> ReportResult:
    return get_report_service().generate(request)


__all__ = [
    "ReportAnalysis",
    "ReportRequest",
    "ReportResult",
    "ReportService",
    "ReservoirInterval",
    "generate_report",
    "get_report_service",
]
