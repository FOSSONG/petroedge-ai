import base64
from io import BytesIO
from openpyxl import load_workbook
import pytest
from app.preprocessing.curve_aliases import canonical_name_for,header_parts
from app.ml_lifecycle.research_excel import workbook_base64

@pytest.mark.parametrize("header,target,unit",[("BOREHOLE-DEPTH.M","depth_m","M"),("GR_COMP.GAPI","gamma_ray_api","GAPI"),("RHOB_COMP.G/C3","density_gcc","G/C3"),("NPHI_COMP.V/V","neutron_porosity_vv","V/V"),("RT_COMP.OHMM","resistivity_ohmm","OHMM"),("RHOB [kg/m3]","density_gcc","KG/M3"),("GR (API)","gamma_ray_api","API")])
def test_explicit_units(header,target,unit):
 assert canonical_name_for(header)==target
 assert header_parts(header)[1]==unit

def test_excel_preserves_rows_and_literal_text():
 result={"unit":"v/v","model_name":"=UNTRUSTED()","samples":[{"source_row":i,"prediction":None if i==1 else .2,"status":"withheld" if i==1 else "predicted","reasons":["=BAD()"] if i==1 else []} for i in range(18373)]}
 wb=load_workbook(BytesIO(base64.b64decode(workbook_base64(result))))
 assert wb["Predictions"].max_row==18374
 assert wb["Predictions"]["B3"].value is None
 assert wb["Predictions"]["E3"].data_type=="s"
 assert wb["Provenance and limitations"]["B3"].value=="=UNTRUSTED()"
 assert wb["Provenance and limitations"]["B3"].data_type=="s"
