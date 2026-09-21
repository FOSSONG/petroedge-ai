"""Excel export of server-computed research results, including withheld rows."""
import base64
import json
from io import BytesIO
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

def workbook_base64(result):
    workbook = Workbook(write_only=True)
    def append(sheet, values):
        cells=[]
        for value in values:
            cell=WriteOnlyCell(sheet, value=value)
            if isinstance(value,str): cell.data_type="s"
            cells.append(cell)
        sheet.append(cells)
    sheet=workbook.create_sheet("Predictions")
    sheet.freeze_panes="A2"
    append(sheet,["Source row (zero-based)", "Prediction", "Unit", "Status", "Reasons"])
    for row in result["samples"]:
        append(sheet,[row["source_row"],row["prediction"],result["unit"],row["status"],"; ".join(row["reasons"])])
    meta=workbook.create_sheet("Provenance and limitations")
    append(meta,["Property","Value"])
    for key,value in result.items():
        if key!="samples": append(meta,[key,json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else value])
    output=BytesIO();workbook.save(output)
    return base64.b64encode(output.getvalue()).decode("ascii")
