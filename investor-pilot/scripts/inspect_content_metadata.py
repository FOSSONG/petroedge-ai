"""Bounded, read-only spreadsheet and SEG-Y metadata inventory."""
from pathlib import Path
import json, re, datetime
ROOT = Path("E:/PETROEDGE_AI/codex")
SOURCE = Path("F:/DATA_PETROEDGE_AI")
OUT = ROOT / "data/catalogs/content_metadata_001"
OUT.mkdir(parents=True, exist_ok=False)
def hints(text):
    patterns = {"pvt":r"\bpvt\b|bubble.?point|formation.volume.factor|fluid.composition",
      "production":r"production|oil.rate|gas.rate|water.rate|cumulative",
      "core":r"core.plug|core.analysis|rcal|scal",
      "porosity":r"porosity|\bphie\b|\bphit\b", "permeability":r"permeability|\bperm\b",
      "saturation":r"saturation|\bswat\b|\bsoil\b|\bsgas\b",
      "trajectory_tops":r"deviation|inclination|azimuth|formation.top",
      "drilling":r"weight.on.bit|standpipe|torque|\brop\b|\bwob\b",
      "seismic":r"seismic|inline|crossline|cdp"}
    return [key for key, pattern in patterns.items() if re.search(pattern,text,re.I)]
items=[]
paths=sorted([p for p in SOURCE.rglob("*") if p.suffix.lower() in {".xls",".xlsx",".sgy",".segy"} and not p.name.startswith("~$")])
for i,path in enumerate(paths):
    record={"source_path":str(path),"relative_path":path.relative_to(SOURCE).as_posix(),"status":"metadata_only_not_training_ready","errors":[]}
    text=""
    try:
        if path.suffix.lower() in {".sgy",".segy"}:
            with path.open("rb") as stream:
                raw=stream.read(3200)
            options=[(encoding,raw.decode(encoding,errors="replace")) for encoding in ["ascii","cp500"]]
            encoding,text=max(options,key=lambda x:sum(c.isprintable() and c!="\ufffd" for c in x[1]))
            record.update({"text_header_encoding_candidate":encoding,"text_header":text,"bytes_read":len(raw),
                "limits":"First 3200 bytes only; binary header, traces, units, CRS and survey overlap not validated"})
        elif path.stat().st_size > 32*1024**2:
            record["errors"].append("Workbook exceeds bounded 32MiB metadata pass")
        elif path.suffix.lower()==".xlsx":
            import openpyxl
            book=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
            record["sheets"]=[]
            try:
                for sheet in book.worksheets:
                    preview=[[str(value)[:240] if value is not None else "" for value in row] for row in sheet.iter_rows(min_row=1,max_row=12,max_col=30,values_only=True)]
                    record["sheets"].append({"name":sheet.title,"reported_rows":sheet.max_row,"reported_columns":sheet.max_column,"first_12_rows":preview})
                text=json.dumps(record["sheets"],ensure_ascii=False)
            finally:
                book.close()
        else:
            import xlrd
            book=xlrd.open_workbook(str(path),on_demand=True)
            record["sheets"]=[]
            try:
                for sheet in book.sheets():
                    preview=[[str(v)[:240] for v in sheet.row_values(row,0,min(30,sheet.ncols))] for row in range(min(12,sheet.nrows))]
                    record["sheets"].append({"name":sheet.name,"reported_rows":sheet.nrows,"reported_columns":sheet.ncols,"first_12_rows":preview})
                text=json.dumps(record["sheets"],ensure_ascii=False)
            finally:
                book.release_resources()
        record["purpose_hints"]=hints(path.name+" "+text)
        record["basin_mentions"]=re.findall(r".{0,70}NIGER[\s_-]+DELTA.{0,70}",text,re.I)[:5]
        record["geography_status"]="requires_source_provenance_review"
    except Exception as exc:
        record["errors"].append(type(exc).__name__+": "+str(exc)[:300])
    items.append(record)
    with (OUT/"metadata.jsonl").open("a",encoding="utf-8") as stream:
        stream.write(json.dumps(record,ensure_ascii=False)+"\n")
    if (i+1)%20==0:
        print(json.dumps({"metadata_files":i+1,"total":len(paths)}),flush=True)
summary={"created_utc":datetime.datetime.now(datetime.timezone.utc).isoformat(),"files_inspected":len(items),
 "workbooks":sum("sheets" in r for r in items),"seismic_text_headers":sum("text_header" in r for r in items),
 "errors":sum(bool(r["errors"]) for r in items),"limits":"Bounded metadata only; hints are not ground-truth labels or region verification"}
(OUT/"SUMMARY.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")
print(json.dumps(summary),flush=True)

