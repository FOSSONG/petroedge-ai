"""Read core workbooks into traceable records; never infer missing units as verified."""
import sys
# xlrd is absent from the bundled runtime; use the installed legacy-XLS reader.
sys.path.append(r"E:\PETROEDGE_AI\codex\runtime\venv\Lib\site-packages")
import collections,hashlib,json,math,re
from pathlib import Path
import openpyxl,xlrd
BASE=Path(r"E:\PETROEDGE_AI\codex\data")
OUT=BASE/"readiness/log_core_001"
def num(v):
    if isinstance(v,bool):return None
    try:n=float(v);return n if math.isfinite(n) and n not in (-999,-999.25,-9999) else None
    except (ValueError,TypeError):return None
def cell(col,row):
    s=""
    while col:col,a=divmod(col-1,26);s=chr(65+a)+s
    return s+str(row)
def val(row,c):return row[c-1] if c and len(row)>=c else ""
def clean(v):
    if v is None:return ""
    if isinstance(v,(str,int,float,bool)):return v
    return str(v)
def sheets(path):
    if path.suffix.lower()==".xlsx":
        b=openpyxl.load_workbook(path,read_only=True,data_only=True)
        result=[{"name":s.title,"rows":[[clean(v) for v in row] for row in s.iter_rows(values_only=True)]} for s in b.worksheets]
        b.close()
    else:
        b=xlrd.open_workbook(path,on_demand=True);result=[]
        for s in b.sheets():
            rows=[]
            for i in range(s.nrows):
                rows.append([xlrd.error_text_from_code.get(s.cell(i,j).value,"#ERROR") if s.cell(i,j).ctype==xlrd.XL_CELL_ERROR else clean(s.cell_value(i,j)) for j in range(s.ncols)])
            result.append({"name":s.name,"rows":rows})
        b.release_resources()
    return result
def extract(book):
    samples=[];shifts=[];coverage=[]
    path=book["source_path"];name=Path(path).name
    for sh in book["sheets"]:
        rows=sh["rows"];sn=sh["name"];spec=None;start=0;well_col=0;static="";depth_col=0;unit="";desc=0;lith=0;samplecol=0;corecol=0;shiftdepth=0;targets=[]
        if name.startswith("GABO-51 PRELIMINARY") and sn.startswith("RR"):
            start=10;static="GABO-51";depth_col=2;unit="m";samplecol=1;desc=9
            targets=[(3,"porosity","%","helium @190 bar"),(4,"permeability","mD","air horizontal @190 bar"),(5,"porosity","%","helium @225 bar"),(6,"permeability","mD","air horizontal @225 bar")]
        elif name.startswith("GABO12-13") and sn=="CoreDB":
            start=7;well_col=2;depth_col=3;samplecol=1
            targets=[(5,"porosity","V/V","source experiment"),(6,"permeability","mD","gas; condition in column G")]
        elif sn=="Summary" and name.startswith("Freeman (1-4ST)"):
            start=6;well_col=1;depth_col=6;shiftdepth=7;samplecol=3;corecol=4;desc=12
            targets=[(c,t,"",f"source {cond}; units require confirmation") for c,t,cond in [(14,"porosity","por_1000"),(15,"permeability","ka_1000"),(16,"permeability","kb_1000"),(17,"porosity","por_3000"),(18,"permeability","ka_3000"),(19,"permeability","kb_3000"),(20,"porosity","por_4500"),(21,"permeability","ka_4500"),(22,"permeability","kb_4500")]]
        elif sn=="Corelab Poroperm" and name.startswith("Freeman-4"):
            start=11;static="Freeman-4";depth_col=2;unit="ft";samplecol=1;desc=7
            targets=[(3,"permeability","mD","air @1000 PSIG"),(4,"permeability","mD","brine @1000 PSIG"),(5,"porosity","%","helium @1000 PSIG")]
        elif sn=="Sheet1" and name=="15919A.XLS" or sn=="A" and name=="15919bt2.xls":
            start=7 if name=="15919A.XLS" else 6;static="15/9-19 A" if name=="15919A.XLS" else "15/9-19 BT2";depth_col=3;unit="m";samplecol=2;corecol=1;desc=17
            targets=[(4,"permeability","mD","gas horizontal"),(6,"permeability","mD","liquid horizontal"),(7,"permeability","mD","gas vertical"),(9,"permeability","mD","liquid vertical"),(10,"porosity","%","horizontal"),(11,"porosity","%","vertical"),(13,"oil_saturation","%","core laboratory; in-situ representativeness unverified"),(14,"water_saturation","%","core laboratory; in-situ representativeness unverified")]
        elif sn=="cores1to9" and "Alwyn - Annex 1" in name:
            start=5;static="ALWYN 3/9A-4";depth_col=3;samplecol=2;corecol=11
            targets=[(4,"permeability","","KA; units not declared"),(5,"permeability","","KL; units not declared"),(6,"porosity","","GEX; units not declared"),(7,"porosity","","FLD; units not declared"),(8,"oil_saturation","","core OIL; units/representativeness unverified"),(9,"water_saturation","","core WTR; units/representativeness unverified")]
        elif sn=="Sdewall_Core_Sample_ Data":
            start=2;well_col=2;depth_col=3;samplecol=5;lith=6;desc=7
        coverage.append({"file_id":book["file_id"],"source_path":path,"sheet":sn,"rows":len(rows),"columns":max(map(len,rows),default=0),"formula_error_cells":sum(isinstance(v,str) and v.startswith(("#DIV/0!","#REF!","#VALUE!","#N/A","#NUM!")) for row in rows for v in row),"extraction":"sample_schema" if start else "retained_for_schema_review"})
        if start:
            well=static
            for ri,row in enumerate(rows,start=1):
                if ri<start:continue
                if name.startswith("Freeman-4") and sn=="Corelab Poroperm" and str(val(row,1)).strip().upper()=="WELL ID":
                    well=str(val(row,2)).strip().lstrip(":").strip().replace(" ST","ST")
                    continue
                if name.startswith("GABO12-13"):
                    if num(val(row,1)) is None:continue
                    if str(val(row,well_col)).strip() and not re.fullmatch(r"GABO-\d+",str(val(row,well_col)).strip(),re.I):continue
                if well_col and str(val(row,well_col)).strip():well=str(val(row,well_col)).strip()
                depth=num(val(row,depth_col))
                if depth is None or depth<=0 or not well:continue
                targetlist=[]
                for col,task,u,condition in targets:
                    raw=val(row,col);v=num(raw)
                    if v is not None:
                        targetlist.append({"task":task,"value":v,"unit":u,"condition":condition,"cell":cell(col,ri),"status":"source_table_value_pending_qualification","raw":raw})
                record={"sample_record_id":hashlib.sha256(f'{book["sha256"]}|{sn}|{ri}'.encode()).hexdigest()[:24],"source_path":path,"source_sha256":book["sha256"],"source_file_id":book["file_id"],"collection":book["collection"],"sheet":sn,"excel_row":ri,"well_source":well,"sample_id_source":str(val(row,samplecol)),"core_id_source":str(val(row,corecol)),"core_depth":depth,"depth_unit":unit,"depth_unit_status":"explicit_source_header" if unit else "unresolved","depth_cell":cell(depth_col,ri),"source_log_depth":num(val(row,shiftdepth)),"targets":targetlist,"lithology_class_source":str(val(row,lith)).strip(),"description_source":str(val(row,desc)).strip(),"training_ready":False}
                if not targetlist and not record["lithology_class_source"] and not record["description_source"]:continue
                samples.append(record)
        if name=="GABO-51_Core_to_log_shift.xls" and sn=="Sheet1":
            for ri,row in enumerate(rows,1):
                if ri>=5 and all(num(val(row,c)) is not None for c in [2,3,4,5]):
                    shifts.append({"well":"GABO-51","core_id":str(val(row,1)),"core_top_m":num(val(row,2)),"core_base_m":num(val(row,3)),"log_top_m":num(val(row,4)),"log_base_m":num(val(row,5)),"source_path":path,"sheet":sn,"excel_row":ri,"status":"source_estimated_shift"})
        elif "core shift" in name.lower():
            for ri,row in enumerate(rows,1):
                if ri>=2 and all(num(val(row,c)) is not None for c in [8,9,10,11]):
                    shifts.append({"well":str(val(row,2)),"core_id":str(val(row,3)),"core_top_m":num(val(row,8)),"core_base_m":num(val(row,9)),"log_top_m":num(val(row,10)),"log_base_m":num(val(row,11)),"source_path":path,"sheet":sn,"excel_row":ri,"status":"source_shift_table"})
    return samples,shifts,coverage
def run():
    OUT.mkdir(parents=True,exist_ok=True);raw=OUT/"core_workbooks";raw.mkdir(exist_ok=True)
    catalog=[json.loads(s) for s in (BASE/"catalogs/all_data_003/catalog.jsonl").read_text(encoding="utf-8").splitlines()]
    inputs=[r for r in catalog if r["detected_format"] in {"XLS","XLSX"} and "core_lab" in r["varieties"]]
    seen={};books=[];dupes=[];samples=[];shifts=[];coverage=[]
    for row in inputs:
        p=Path(row["source_path"])
        with p.open("rb") as f:sha=hashlib.file_digest(f,"sha256").hexdigest()
        if sha in seen:dupes.append({"source_path":str(p),"sha256":sha,"same_as":seen[sha]});continue
        seen[sha]=str(p)
        book={k:row[k] for k in ["file_id","source_path","collection"]};book["sha256"]=sha
        try:book["sheets"]=sheets(p);book["status"]="read";a,b,c=extract(book);samples.extend(a);shifts.extend(b);coverage.extend(c)
        except Exception as e:book.update(sheets=[],status="error",error=str(e))
        (raw/(row["file_id"]+".json")).write_text(json.dumps(book,ensure_ascii=True),encoding="utf-8")
        books.append({k:v for k,v in book.items() if k!="sheets"})
        print("Core workbook",len(books),p.name,book["status"],flush=True)
    for name,data in [("core_samples",samples),("core_shifts",shifts),("core_sheet_coverage",coverage),("core_workbook_inventory",books),("core_duplicate_files",dupes)]:
        (OUT/(name+".jsonl")).write_text("".join(json.dumps(r,ensure_ascii=True)+"\n" for r in data),encoding="utf-8")
    summary={"candidate_workbooks":len(inputs),"unique_workbooks":len(books),"exact_duplicate_files":len(dupes),"read_errors":[r for r in books if r["status"]=="error"],"extracted_source_rows":len(samples),"shift_records":len(shifts),"source_well_names":sorted({r["well_source"] for r in samples}),"formula_error_cells":sum(r["formula_error_cells"] for r in coverage)}
    (OUT/"CORE_SUMMARY.json").write_text(json.dumps(summary,indent=2),encoding="utf-8");print(json.dumps(summary),flush=True)
if __name__=="__main__":run()
